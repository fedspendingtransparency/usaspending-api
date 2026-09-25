import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from usaspending_api.llm.evals.exceptions import EvalError
from usaspending_api.llm.evals.models import EvalCase, EvalResult, EvalSummary

SUMMARY_FIELDS = (
    "assistant",
    "dataset",
    "case_count",
    "total_case_count",
    "unrun_count",
    "passed_count",
    "failed_count",
    "score",
    "fail_under",
    "passed",
)
CASE_FIELDS = (
    "case_name",
    "case_id",
    "query",
    "status",
    "score",
    "tool_passed",
    "tool_score",
    "tool_message",
    "expected_tools",
    "actual_tools",
    "output_passed",
    "output_score",
    "output_message",
    "expected_output",
    "actual_output",
    "case_metadata",
    "execution_metadata",
    "assistant",
    "assistant_id",
    "ai_model_id",
    "system_prompt_id",
    "inference_config_temp",
    "inference_config_top_p",
    "inference_config_max_tokens",
    "inference_config_stop_sequences",
    "input_tokens",
    "output_tokens",
    "latencyMs",
)
TAG_FIELDS = (
    "tag",
    "total_cases",
    "run_cases",
    "not_run_cases",
    "passed_count",
    "failed_count",
    "average_score",
    "pass_rate",
)
# Generic string placed in cells where a value is unavailable/not applicable.
NOT_RUN = "N/A"


def match_dict(match: Any) -> dict[str, Any] | None:
    """Converts one MatchResult to JSON-compatible data."""
    if match is None:
        return None

    return {
        "passed": match.passed,
        "score": match.score,
        "message": match.message,
        "expected": match.expected,
        "actual": match.actual,
    }


def _execution_fields(metadata: dict[str, Any]) -> dict[str, Any]:
    return {
        "assistant": metadata.get("assistant", NOT_RUN),
        "assistant_id": metadata.get("assistant_id", NOT_RUN),
        "ai_model_id": metadata.get("ai_model_id", NOT_RUN),
        "system_prompt_id": metadata.get("system_prompt_id", NOT_RUN),
        "inference_config_temp": metadata.get("inference_config_temp", NOT_RUN),
        "inference_config_top_p": metadata.get("inference_config_top_p", NOT_RUN),
        "inference_config_max_tokens": metadata.get("inference_config_max_tokens", NOT_RUN),
        "inference_config_stop_sequences": metadata.get(
            "inference_config_stop_sequences",
            NOT_RUN,
        ),
        "input_tokens": metadata.get("input_tokens", NOT_RUN),
        "output_tokens": metadata.get("output_tokens", NOT_RUN),
        "latencyMs": metadata.get("latencyMs", NOT_RUN),
    }


def result_to_dict(result: EvalResult) -> dict[str, Any]:
    """Converts one executed case result to a detailed evaluation row."""
    # Handle execution errors
    if result.error:
        return {
            "case_name": result.case_name,
            "case_id": result.case_name,
            "query": result.details.get("case_input", {}).get("query", NOT_RUN),
            "status": "ERROR",
            "passed": False,
            "score": 0.0,
            "tool_call_match": NOT_RUN,
            "output_match": NOT_RUN,
            "tool_passed": NOT_RUN,
            "tool_score": NOT_RUN,
            "tool_message": result.error,
            "expected_tools": NOT_RUN,
            "actual_tools": NOT_RUN,
            "output_passed": NOT_RUN,
            "output_score": NOT_RUN,
            "output_message": result.error,
            "expected_output": NOT_RUN,
            "actual_output": NOT_RUN,
            "case_metadata": result.details.get("case_metadata", {}),
            "execution_metadata": NOT_RUN,
            "assistant": NOT_RUN,
            "assistant_id": NOT_RUN,
            "ai_model_id": NOT_RUN,
            "system_prompt_id": NOT_RUN,
            "inference_config_temp": NOT_RUN,
            "inference_config_top_p": NOT_RUN,
            "inference_config_max_tokens": NOT_RUN,
            "inference_config_stop_sequences": NOT_RUN,
            "input_tokens": NOT_RUN,
            "output_tokens": NOT_RUN,
            "latencyMs": NOT_RUN,
        }

    tool_match = match_dict(result.tool_call_match)
    output_match = match_dict(result.output_match)
    case_metadata = result.details.get("case_metadata", {})
    execution_metadata = result.details.get("execution_metadata", {})
    case_input = result.details.get("case_input", {})

    return {
        "case_name": result.case_name,
        "case_id": result.case_name,
        "query": case_input.get("query", NOT_RUN),
        "status": "PASS" if result.passed else "FAIL",
        "passed": result.passed,
        "score": result.score,
        "tool_call_match": tool_match,
        "output_match": output_match,
        "tool_passed": tool_match["passed"] if tool_match else NOT_RUN,
        "tool_score": tool_match["score"] if tool_match else NOT_RUN,
        "tool_message": tool_match["message"] if tool_match else NOT_RUN,
        "expected_tools": tool_match["expected"] if tool_match else NOT_RUN,
        "actual_tools": tool_match["actual"] if tool_match else NOT_RUN,
        "output_passed": output_match["passed"] if output_match else NOT_RUN,
        "output_score": output_match["score"] if output_match else NOT_RUN,
        "output_message": output_match["message"] if output_match else NOT_RUN,
        "expected_output": output_match["expected"] if output_match else NOT_RUN,
        "actual_output": output_match["actual"] if output_match else NOT_RUN,
        "case_metadata": case_metadata,
        "execution_metadata": execution_metadata,
        **_execution_fields(execution_metadata),
    }


def unrun_case_to_dict(summary: EvalSummary, case: EvalCase) -> dict[str, Any]:
    """Converts an excluded case into a row with explicit N/A actual fields."""
    expected_tools = list(case.expected_tool_calls)
    reason = summary.unrun_reasons.get(case.name, "case was not run")

    return {
        "case_name": case.name,
        "case_id": case.name,
        "query": case.input.get("query", ""),
        "status": "NOT RUN",
        "passed": None,
        "score": None,
        "tool_call_match": NOT_RUN,
        "output_match": NOT_RUN,
        "tool_passed": NOT_RUN,
        "tool_score": NOT_RUN,
        "tool_message": reason,
        "expected_tools": expected_tools,
        "actual_tools": NOT_RUN,
        "output_passed": NOT_RUN,
        "output_score": NOT_RUN,
        "output_message": reason,
        "expected_output": case.expected_output,
        "actual_output": NOT_RUN,
        "case_metadata": case.metadata,
        "execution_metadata": NOT_RUN,
        "assistant": summary.assistant,
        "assistant_id": NOT_RUN,
        "ai_model_id": NOT_RUN,
        "system_prompt_id": NOT_RUN,
        "inference_config_temp": NOT_RUN,
        "inference_config_top_p": NOT_RUN,
        "inference_config_max_tokens": NOT_RUN,
        "inference_config_stop_sequences": NOT_RUN,
        "input_tokens": NOT_RUN,
        "output_tokens": NOT_RUN,
        "latencyMs": NOT_RUN,
    }


def case_rows(summary: EvalSummary) -> list[dict[str, Any]]:
    """Returns executed and excluded cases in one report-ready collection."""
    executed = [result_to_dict(result) for result in summary.results]
    unrun = [unrun_case_to_dict(summary, case) for case in summary.unrun_cases]
    return executed + unrun


def summary_to_dict(summary: EvalSummary) -> dict[str, Any]:
    """Converts aggregate run results to a machine-readable summary."""
    return {
        "assistant": summary.assistant,
        "dataset": summary.dataset,
        "case_count": summary.case_count,
        "total_case_count": summary.total_case_count,
        "unrun_count": summary.unrun_count,
        "passed_count": summary.passed_count,
        "failed_count": summary.failed_count,
        "score": summary.score,
        "fail_under": summary.fail_under,
        "passed": summary.passed,
    }


def tag_rows(summary: EvalSummary, rows: list[dict[str, Any]] | None = None) -> list[dict[str, Any]]:
    """Aggregates total, executed, unrun, passed, failed, and score by tag."""
    rows = rows if rows is not None else case_rows(summary)
    grouped: dict[str, list[dict[str, Any]]] = defaultdict(list)

    for row in rows:
        tags = row.get("case_metadata", {}).get("tags", [])
        for tag in tags or ["(untagged)"]:
            grouped[tag].append(row)

    results = []
    for tag, tag_cases in sorted(grouped.items()):
        executed = [case for case in tag_cases if case["status"] in {"PASS", "FAIL", "ERROR"}]
        passed = sum(case["status"] == "PASS" for case in executed)
        scores = [case["score"] for case in executed if isinstance(case["score"], (int, float))]
        results.append(
            {
                "tag": tag,
                "total_cases": len(tag_cases),
                "run_cases": len(executed),
                "not_run_cases": len(tag_cases) - len(executed),
                "passed_count": passed,
                "failed_count": len(executed) - passed,
                "average_score": sum(scores) / len(scores) if scores else None,
                "pass_rate": passed / len(executed) if executed else None,
            }
        )

    return results


def report_to_dict(summary: EvalSummary) -> dict[str, Any]:
    """Builds the complete summary, tag, and row-level export payload."""
    summary_data = summary_to_dict(summary)
    rows = case_rows(summary)
    tags = tag_rows(summary, rows)

    return {
        **summary_data,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "summary": summary_data,
        "tags": tags,
        "cases": rows,
        "results": rows,
    }


def render_json(summary: EvalSummary) -> str:
    """Returns the complete summary and case analysis as formatted JSON."""
    return json.dumps(report_to_dict(summary), indent=2, sort_keys=True, default=str)


def _display_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(value, sort_keys=True, default=str)
    return str(value)


def _style_sheet(sheet: Any, header_fill: Any, header_font: Any, get_column_letter: Any) -> None:
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for column_cells in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column_cells) + 2, 60)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = width


def write_xlsx(summary: EvalSummary, output_path: Path) -> None:
    """Exports Summary, Tags, and row-level Evaluation worksheets."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise EvalError("XLSX export requires openpyxl.") from error

    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "Summary"
    tags_sheet = workbook.create_sheet("Tags")
    evaluation_sheet = workbook.create_sheet("Evaluation")
    rows = case_rows(summary)

    summary_data = summary_to_dict(summary)
    summary_sheet.append(["Metric", "Value"])
    summary_sheet.append(["Generated At", datetime.now(timezone.utc).isoformat()])
    for field in SUMMARY_FIELDS:
        summary_sheet.append([field, _display_value(summary_data[field])])

    tags_sheet.append(list(TAG_FIELDS))
    for tag in tag_rows(summary, rows):
        tags_sheet.append([_display_value(tag.get(field)) for field in TAG_FIELDS])

    evaluation_sheet.append(list(CASE_FIELDS))
    for row in rows:
        evaluation_sheet.append([_display_value(row.get(field)) for field in CASE_FIELDS])

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    for sheet in (summary_sheet, tags_sheet, evaluation_sheet):
        _style_sheet(sheet, header_fill, header_font, get_column_letter)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    except OSError as error:
        raise EvalError(f"Unable to write XLSX evaluation report '{output_path}': {error}") from error


def render_text(summary: EvalSummary) -> str:
    """Returns a concise report intended for terminal use and CI logs."""
    lines = [
        f"Assistant: {summary.assistant}",
        f"Dataset: {summary.dataset}",
        f"Cases run: {summary.case_count}",
        f"Total cases: {summary.total_case_count}",
        f"Not run: {summary.unrun_count}",
        f"Passed: {summary.passed_count}",
        f"Failed: {summary.failed_count}",
        f"Score: {summary.score:.2%}",
    ]

    if summary.fail_under is not None:
        lines.append(f"Required score: {summary.fail_under:.2%}")

    lines.extend([f"Result: {'PASS' if summary.passed else 'FAIL'}", ""])

    for row in case_rows(summary):
        lines.append(f"{row['status']}  {row['case_name']}  ({_display_value(row['score'])})")
        if row["status"] == "ERROR":
            lines.append(f"  Error: {row['tool_message']}")
        elif row["status"] == "FAIL":
            if row["tool_passed"] is False:
                lines.append(f"  Tool calls: {row['tool_message']}")
            if row["output_passed"] is False:
                lines.append(f"  Output: {row['output_message']}")
        elif row["status"] == "NOT RUN":
            lines.append(f"  Reason: {row['tool_message']}")

    return "\n".join(lines)
