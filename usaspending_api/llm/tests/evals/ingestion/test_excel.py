import io
import json

import openpyxl
import pytest

from usaspending_api.llm.evals.exceptions import DatasetError
from usaspending_api.llm.evals.ingestion.excel import read_workbook
from usaspending_api.llm.evals.ingestion.models import GroundTruthSource


def workbook_bytes(
    ground_truth_headers=None,
    ground_truth_rows=None,
    dictionary_headers=None,
    dictionary_rows=None,
):
    workbook = openpyxl.Workbook()
    ground_truth = workbook.active
    ground_truth.title = "Ground Truth"
    dictionary = workbook.create_sheet("Filter Dictionary")

    ground_truth.append(
        ground_truth_headers
        or [
            "id",
            "query",
            "expected_output",
            "expected_tools",
            "tags",
            "notes",
            "approved",
            "sme_validation_notes",
        ]
    )
    for row in ground_truth_rows or [
        [
            1,
            "Find Clark Construction contracts",
            json.dumps({"recipient": "Clark Construction"}),
            json.dumps(["lookup recipient", "execute filter"]),
            json.dumps(["recipient"]),
            "",
            True,
            "",
        ]
    ]:
        ground_truth.append(row)

    dictionary.append(
        dictionary_headers
        or ["Filter", "Subfilter", "Naming Convention", "Backend Naming Convention"]
    )
    for row in dictionary_rows or [
        ["Filters", "Recipient", "recipient", "selectedRecipients"],
        ["Filters", "Time Period", "time_period.fiscal_year", "timePeriodFY"],
        ["Tools", "Recipient", "lookup recipient", "lookup_recipient"],
        ["Tools", "Completion", "execute filter", "execute_filter"],
    ]:
        dictionary.append(row)

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def source(content: bytes) -> GroundTruthSource:
    return GroundTruthSource(
        content=content,
        source_name="fixture.xlsx",
        source_reference="fixture.xlsx",
    )


def test_read_workbook_reads_actual_dictionary_headers_and_mappings():
    result = read_workbook(source(workbook_bytes()))

    assert len(result.ground_truth_rows) == 1
    assert result.ground_truth_rows[0]["id"] == 1
    assert len(result.filter_mappings) == 4
    assert result.filter_mappings[0].source_name == "recipient"
    assert result.filter_mappings[0].target_name == "selectedRecipients"
    assert result.filter_mappings[0].filter_name == "Filters"
    assert result.filter_mappings[0].subfilter == "Recipient"


def test_read_workbook_rejects_missing_sheet():
    workbook = openpyxl.Workbook()
    workbook.active.title = "Ground Truth"
    stream = io.BytesIO()
    workbook.save(stream)

    with pytest.raises(DatasetError, match="missing worksheet"):
        read_workbook(source(stream.getvalue()))


def test_read_workbook_rejects_missing_column():
    with pytest.raises(DatasetError, match="missing column"):
        read_workbook(
            source(
                workbook_bytes(
                    ground_truth_headers=["id"],
                    ground_truth_rows=[[1]],
                )
            )
        )


def test_read_workbook_rejects_duplicate_naming_convention():
    duplicate_mapping = [
        ["Filters", "Recipient", "recipient", "selectedRecipients"],
        ["Filters", "Recipient Alias", "recipient", "selectedRecipients"],
    ]

    with pytest.raises(DatasetError, match="duplicate Naming Convention"):
        read_workbook(source(workbook_bytes(dictionary_rows=duplicate_mapping)))
