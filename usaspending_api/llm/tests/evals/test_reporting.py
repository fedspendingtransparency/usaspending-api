import json

from openpyxl import load_workbook

from usaspending_api.llm.evals.models import EvalCase, EvalResult, EvalSummary, MatchResult
from usaspending_api.llm.evals.reporting import report_to_dict, write_xlsx


def summary():
    tool_match = MatchResult(
        passed=True,
        score=1.0,
        expected=[{"name": "execute_filter"}],
        actual=[{"name": "execute_filter"}],
        message="Tool call sequence matches expected values.",
    )
    output_match = MatchResult(
        passed=False,
        score=0.0,
        expected={"timePeriodFY": ["2025"]},
        actual={"timePeriodFY": ["2024"]},
        message="Output differs at: timePeriodFY",
    )

    return EvalSummary(
        assistant="filter_search",
        dataset="ground_truth",
        results=(
            EvalResult(
                case_name="1",
                passed=False,
                score=0.5,
                tool_call_match=tool_match,
                output_match=output_match,
                details={
                    "case_metadata": {
                        "tags": ["temporal"],
                        "notes": "Example case",
                    },
                    "execution_metadata": {
                        "session_id": "123",
                    },
                },
            ),
        ),
        score=0.5,
        passed=False,
        fail_under=1.0,
        unrun_cases=(
            EvalCase(
                name="2",
                input={"query": "Unrun query"},
                expected_output={"timePeriodFY": ["2025"]},
                metadata={"tags": ["temporal"], "approved": False},
            ),
        ),
        unrun_reasons={"2": "case is not approved"},
    )


def test_report_to_dict_contains_summary_and_case_rows():
    report = report_to_dict(summary())

    assert report["summary"]["score"] == 0.5
    assert report["summary"]["failed_count"] == 1
    assert report["summary"]["total_case_count"] == 2
    assert report["summary"]["unrun_count"] == 1
    assert len(report["cases"]) == 2
    assert report["cases"][0]["case_name"] == "1"
    assert report["cases"][0]["status"] == "FAIL"
    assert report["cases"][0]["output_passed"] is False
    assert report["cases"][0]["expected_output"] == {"timePeriodFY": ["2025"]}
    assert report["cases"][0]["actual_output"] == {"timePeriodFY": ["2024"]}
    assert report["cases"][1]["case_name"] == "2"
    assert report["cases"][1]["status"] == "NOT RUN"
    assert report["cases"][1]["actual_output"] == "N/A"
    assert report["cases"][1]["tool_message"] == "case is not approved"
    assert report["tags"] == [
        {
            "tag": "temporal",
            "total_cases": 2,
            "run_cases": 1,
            "not_run_cases": 1,
            "passed_count": 0,
            "failed_count": 1,
            "average_score": 0.5,
            "pass_rate": 0.0,
        }
    ]

    # The top-level fields and legacy results remain available for existing JSON
    # consumers while new consumers can use summary/cases explicitly.
    assert report["assistant"] == "filter_search"
    assert report["results"] == report["cases"]
    json.dumps(report)


def test_write_xlsx_creates_summary_and_case_results_sheets(tmp_path):
    output_path = tmp_path / "evaluation_report.xlsx"

    write_xlsx(summary(), output_path)

    workbook = load_workbook(output_path, read_only=True, data_only=True)

    assert workbook.sheetnames == ["Summary", "Tags", "Evaluation"]
    assert workbook["Summary"]["A1"].value == "Metric"
    assert workbook["Tags"]["A1"].value == "tag"
    assert workbook["Tags"]["A2"].value == "temporal"
    assert workbook["Evaluation"]["A1"].value == "case_name"
    assert workbook["Evaluation"]["A2"].value == "1"
    assert workbook["Evaluation"]["D2"].value == "FAIL"
    assert workbook["Evaluation"]["A3"].value == "2"
    assert workbook["Evaluation"]["D3"].value == "NOT RUN"
    assert workbook["Evaluation"]["O3"].value == "N/A"
